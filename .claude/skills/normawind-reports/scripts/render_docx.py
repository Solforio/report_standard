# -*- coding: utf-8 -*-
"""Renderer DOCX editable. Estilos reales de Word (Heading 1-3, Caption, List Bullet) + índice automático."""
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor, Emu
from PIL import Image
import nwcommon as C
from nwcommon import label, section_title, numbered, fmt_date

TEXT_W = 16.6  # cm
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def rgbc(h):
    return RGBColor.from_string(h)


# ---------- helpers OXML ----------
def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    for e in tcPr.findall(qn("w:shd")):
        tcPr.remove(e)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def cell_borders(cell, **kw):
    """kw: top/bottom/left/right = (size_eighths, color) o None"""
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders"); tcPr.append(b)
    for side in ("top", "left", "bottom", "right"):
        v = kw.get(side, "keep")
        if v == "keep":
            continue
        el = b.find(qn("w:" + side))
        if el is None:
            el = OxmlElement("w:" + side); b.append(el)
        if v is None:
            el.set(qn("w:val"), "nil")
        else:
            el.set(qn("w:val"), "single"); el.set(qn("w:sz"), str(v[0])); el.set(qn("w:space"), "0"); el.set(qn("w:color"), v[1])


def cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    tcPr.append(m)


def table_no_borders(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for s in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + s); e.set(qn("w:val"), "nil"); b.append(e)
    tblPr.append(b)


def set_col_widths(tbl, widths_cm):
    tbl.autofit = False
    tblPr = tbl._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = tbl._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths_cm[i] * 567)))
    for row in tbl.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(widths_cm[i])


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit"); e.set(qn("w:val"), "true"); trPr.append(e)


def para_border_bottom(p, color, sz=6):
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:pBdr")
    e = OxmlElement("w:bottom")
    e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "4"); e.set(qn("w:color"), color)
    b.append(e); pPr.append(b)


def right_tab_only(p):
    """Anula las tabulaciones centro/derecha del estilo Header/Footer y fija una a la derecha del texto."""
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    for pos in (4680, 9360):
        t = OxmlElement("w:tab"); t.set(qn("w:val"), "clear"); t.set(qn("w:pos"), str(pos)); tabs.append(t)
    t = OxmlElement("w:tab"); t.set(qn("w:val"), "right"); t.set(qn("w:pos"), str(int(TEXT_W * 567))); tabs.append(t)
    pPr.append(tabs)


def add_field(par, instr, placeholder="1", size=None, color=None):
    def r():
        run = par.add_run()
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = rgbc(color)
        return run
    a = r(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "begin"); a._r.append(f)
    b = r(); i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = " %s " % instr; b._r.append(i)
    c = r(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "separate"); c._r.append(f)
    d = r(); d.text = placeholder
    e = r(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); e._r.append(f)


def style_font(st, name=C.FONT, size=None, color=None, bold=None, italic=None):
    st.font.name = name
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size: st.font.size = Pt(size)
    if color: st.font.color.rgb = rgbc(color)
    if bold is not None: st.font.bold = bold
    if italic is not None: st.font.italic = italic


def rich(par, text, size=None, color=None, bold=None):
    """Texto con **negrita** y *cursiva*."""
    for tok in INLINE.split(str(text)):
        if not tok:
            continue
        if tok.startswith("**"):
            run = par.add_run(tok[2:-2]); run.bold = True
        elif tok.startswith("*"):
            run = par.add_run(tok[1:-1]); run.italic = True
        else:
            run = par.add_run(tok)
            if bold: run.bold = True
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = rgbc(color)


def is_num(s):
    return bool(re.match(r"^[\s<>≈~±+\-−]*\d[\d.,\s%€$]*(\s?[A-Za-zµ°/%²³€.]{0,8})?$", str(s))) and len(str(s)) < 22


# ---------- documento ----------
class Doc:
    def __init__(self, content):
        self.c = content
        self.meta = content["meta"]
        self.lang = self.meta["lang"]
        self.d = Document()
        self.fig_n = 0
        self.tab_n = 0
        self._setup()

    def L(self, key):
        return label(self.lang, key)

    def _setup(self):
        d = self.d
        s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin, s.bottom_margin = Cm(2.9), Cm(2.3)
        s.header_distance, s.footer_distance = Cm(1.0), Cm(0.9)
        s.different_first_page_header_footer = True
        st = d.styles
        style_font(st["Normal"], size=10.5, color=C.INK)
        st["Normal"].paragraph_format.space_after = Pt(6)
        st["Normal"].paragraph_format.line_spacing = 1.12
        for n, sz, sp in (("Heading 1", 16, 18), ("Heading 2", 13, 12), ("Heading 3", 11, 8)):
            style_font(st[n], size=sz, color=C.NAVY if n != "Heading 3" else C.BLUE, bold=True)
            st[n].paragraph_format.space_before = Pt(sp)
            st[n].paragraph_format.space_after = Pt(6)
            st[n].paragraph_format.keep_with_next = True
        style_font(st["Caption"], size=9, color=C.GREY, italic=True, bold=False)
        st["Caption"].paragraph_format.space_after = Pt(10)
        for n in ("List Bullet", "List Number"):
            style_font(st[n], size=10.5, color=C.INK)
            st[n].paragraph_format.space_after = Pt(3)
        for n in ("TOC 1", "TOC 2"):
            try:
                style_font(st[n], size=10.5 if n == "TOC 1" else 10, color=C.INK, bold=(n == "TOC 1"))
            except KeyError:
                pass
        self._header_footer(s)
        d.core_properties.title = self.meta.get("title", "")
        d.core_properties.author = "; ".join(self.meta.get("authors", [])) or "NORMAWIND S.L."
        d.core_properties.language = self.lang

    # ---- cabecera / pie ----
    def _header_footer(self, s):
        m = self.meta
        hp = s.header.paragraphs[0]
        para_border_bottom(hp, C.BLUE, 8)
        right_tab_only(hp)
        hp.add_run().add_picture(C.logo("normawind-horizontal-positive.png"), height=Cm(0.95))
        r = hp.add_run("\t" + C.doc_short_title(m)); r.font.size = Pt(8.5); r.font.color.rgb = rgbc(C.GREY); r.italic = True
        if m.get("example"):
            hp2 = s.header.add_paragraph()
            hp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = hp2.add_run(self.L("example")); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = rgbc(C.BAD)
        elif m.get("draft"):
            hp2 = s.header.add_paragraph(); hp2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = hp2.add_run(self.L("draft")); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = rgbc(C.BAD)
        fp = s.footer.paragraphs[0]
        right_tab_only(fp)
        left = " · ".join(x for x in (m.get("code"), C.doc_short_title(m)) if x)
        r = fp.add_run(left); r.font.size = Pt(8); r.font.color.rgb = rgbc(C.GREY)
        r = fp.add_run("\t%s " % self.L("page")); r.font.size = Pt(8); r.font.color.rgb = rgbc(C.GREY)
        add_field(fp, "PAGE", "1", 8, C.GREY)
        r = fp.add_run(" %s " % self.L("of")); r.font.size = Pt(8); r.font.color.rgb = rgbc(C.GREY)
        add_field(fp, "NUMPAGES", "1", 8, C.GREY)
        if m.get("confidential", True):
            p2 = s.footer.add_paragraph(); r = p2.add_run(self.L("confidential")); r.font.size = Pt(7); r.font.color.rgb = rgbc(C.TAUPE)

    # ---- portada ----
    def cover(self):
        d, m = self.d, self.meta
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(C.logo("normawind-horizontal-positive.png"), width=Cm(6.2))
        sp = d.add_paragraph(); sp.paragraph_format.space_before = Pt(70)
        t = d.add_table(rows=1, cols=1); table_no_borders(t); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_col_widths(t, [TEXT_W])
        c = t.cell(0, 0); shade(c, C.NAVY); cell_margins(c, 500, 500, 500, 500)
        cell_borders(c, left=(48, C.SKY))
        p = c.paragraphs[0]
        r = p.add_run(C.type_name(self.lang, m.get("doc_type", "")).upper()); r.font.size = Pt(9.5); r.bold = True; r.font.color.rgb = rgbc(C.SKY)
        p.paragraph_format.space_after = Pt(10)
        p = c.add_paragraph(); r = p.add_run(m["title"]); r.font.size = Pt(26); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(8)
        if m.get("subtitle"):
            p = c.add_paragraph(); r = p.add_run(m["subtitle"]); r.font.size = Pt(13); r.font.color.rgb = rgbc(C.BEIGE)
        sp = d.add_paragraph(); sp.paragraph_format.space_before = Pt(28)
        rows = [(self.L("client"), m.get("client")), (self.L("project"), m.get("project")),
                (self.L("location"), m.get("location")), (self.L("reference"), m.get("code")),
                (self.L("version"), m.get("version")), (self.L("date"), fmt_date(m.get("date"), self.lang))]
        rows = [r for r in rows if r[1]]
        t = d.add_table(rows=len(rows), cols=2); table_no_borders(t); set_col_widths(t, [4.0, TEXT_W - 4.0])
        for i, (k, v) in enumerate(rows):
            a, b = t.cell(i, 0), t.cell(i, 1)
            ra = a.paragraphs[0].add_run(k.upper()); ra.font.size = Pt(8.5); ra.bold = True; ra.font.color.rgb = rgbc(C.TAUPE)
            rb = b.paragraphs[0].add_run(str(v)); rb.font.size = Pt(11); rb.font.color.rgb = rgbc(C.INK)
            for c_ in (a, b):
                cell_margins(c_, 70, 70, 0, 100); cell_borders(c_, bottom=(4, "D9DEE5"))
        if m.get("client_logo"):
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(24)
            p.add_run().add_picture(C._abs(self.c, m["client_logo"]), height=Cm(1.6))
        if m.get("example"):
            p = d.add_paragraph(); p.paragraph_format.space_before = Pt(18)
            r = p.add_run(self.L("example")); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = rgbc(C.BAD)
        d.add_page_break()

    # ---- control de revisiones ----
    def control_page(self):
        d, m = self.d, self.meta
        p = d.add_paragraph(); r = p.add_run(self.L("revision_history")); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = rgbc(C.NAVY)
        p.paragraph_format.space_after = Pt(10)
        revs = m["revisions"]
        self._simple_table([self.L("version"), self.L("date"), self.L("changes")],
                           [[r_["version"], fmt_date(r_["date"], self.lang), r_["changes"]] for r_ in revs], [2.5, 3.0, TEXT_W - 5.5])
        d.add_paragraph()
        self._simple_table([self.L("authors"), self.L("validated_by")],
                           [["\n".join(m.get("authors", [])) or "-", m.get("validated_by", "-")]], [TEXT_W / 2, TEXT_W / 2])
        d.add_paragraph()
        info = [[self.L("document_name"), m.get("doc_name") or m["title"]], [self.L("reference"), m.get("code", "-")],
                [self.L("version"), m.get("version")], [self.L("date"), fmt_date(m.get("date"), self.lang)]]
        t = d.add_table(rows=len(info), cols=2); table_no_borders(t); set_col_widths(t, [4.2, TEXT_W - 4.2])
        for i, (k, v) in enumerate(info):
            a, b = t.cell(i, 0), t.cell(i, 1)
            shade(a, C.PALEBLUE)
            ra = a.paragraphs[0].add_run(k); ra.bold = True; ra.font.size = Pt(9.5); ra.font.color.rgb = rgbc(C.NAVY)
            rb = b.paragraphs[0].add_run(str(v)); rb.font.size = Pt(9.5)
            for c_ in (a, b):
                cell_margins(c_, 70, 70, 110, 110); cell_borders(c_, bottom=(4, "D9DEE5"))
        d.add_page_break()

    def toc_page(self):
        d = self.d
        p = d.add_paragraph(); r = p.add_run(self.L("contents")); r.bold = True; r.font.size = Pt(16); r.font.color.rgb = rgbc(C.NAVY)
        p.paragraph_format.space_after = Pt(10)
        p = d.add_paragraph()
        add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "…")
        d.add_page_break()

    # ---- bloques ----
    def heading(self, text, level):
        self.d.add_heading(text, level=level)

    def block(self, b):
        d = self.d
        if "p" in b:
            p = d.add_paragraph(); rich(p, b["p"])
        elif "bullets" in b:
            for it in b["bullets"]:
                p = d.add_paragraph(style="List Bullet"); rich(p, it)
        elif "numbered" in b:
            for i, it in enumerate(b["numbered"], 1):  # numeración manual: cada lista reinicia en 1
                p = d.add_paragraph(); pf = p.paragraph_format
                pf.left_indent = Cm(0.9); pf.first_line_indent = Cm(-0.6); pf.space_after = Pt(3)
                pf.tab_stops.add_tab_stop(Cm(0.9))
                r = p.add_run("%d.\t" % i); r.bold = True; r.font.color.rgb = rgbc(C.BLUE)
                rich(p, it)
        elif "table" in b:
            self.table(b["table"])
        elif "keyvalue" in b:
            self.keyvalue(b["keyvalue"])
        elif "figure" in b:
            self.figure(b["figure"])
        elif "kpis" in b:
            self.kpis(b["kpis"])
        elif "callout" in b:
            self.callout(b["callout"])
        elif "pagebreak" in b:
            d.add_page_break()
        elif "subheading" in b:
            d.add_heading(b["subheading"], level=3)

    def _simple_table(self, header, rows, widths):
        t = self.d.add_table(rows=1 + len(rows), cols=len(header)); table_no_borders(t); set_col_widths(t, widths)
        for j, h in enumerate(header):
            c = t.cell(0, j); shade(c, C.NAVY); cell_margins(c, 70, 70, 110, 110)
            r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(255, 255, 255)
        for i, row in enumerate(rows, 1):
            for j, v in enumerate(row):
                c = t.cell(i, j); cell_margins(c, 70, 70, 110, 110); cell_borders(c, bottom=(4, "D9DEE5"))
                lines = str(v).split("\n")
                c.paragraphs[0].add_run(lines[0]).font.size = Pt(9.5)
                for ln in lines[1:]:
                    c.add_paragraph().add_run(ln).font.size = Pt(9.5)
        return t

    def table(self, spec):
        d = self.d
        header, rows = spec["header"], spec["rows"]
        self.tab_n += 1
        if spec.get("caption"):
            cp = d.add_paragraph(style="Caption"); cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_after = Pt(3)
            cp.add_run("%s %d. %s" % (self.L("table"), self.tab_n, spec["caption"]))
        n = len(header)
        w = spec.get("widths") or [1] * n
        tot = float(sum(w)); widths = [TEXT_W * x / tot for x in w]
        t = d.add_table(rows=1 + len(rows), cols=n); table_no_borders(t); set_col_widths(t, widths)
        repeat_header(t.rows[0])
        for j, h in enumerate(header):
            c = t.cell(0, j); shade(c, C.NAVY); cell_margins(c, 60, 60, 90, 90)
            p = c.paragraphs[0]; r = p.add_run(str(h)); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor(255, 255, 255)
            p.paragraph_format.space_after = Pt(0)
            if j > 0 and all(is_num(rw[j]) for rw in rows):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for i, row in enumerate(rows, 1):
            no_split(t.rows[i])
            bold_row = spec.get("total_row") and i == len(rows)
            for j, v in enumerate(row):
                c = t.cell(i, j); cell_margins(c, 50, 50, 90, 90)
                if i % 2 == 0: shade(c, C.LIGHT)
                if bold_row: shade(c, C.PALEBLUE)
                cell_borders(c, bottom=(4, "D9DEE5"))
                p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
                rich(p, v, size=9, bold=bool(bold_row))
                if j > 0 and all(is_num(rw[j]) for rw in rows):
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        keep = len(rows) <= 18  # tablas cortas: no partirlas entre páginas; largas: solo cabecera con la 1ª fila
        for i, row_ in enumerate(t.rows):
            if (keep and i < len(t.rows) - 1) or i == 0:
                for c_ in row_.cells:
                    for p_ in c_.paragraphs:
                        p_.paragraph_format.keep_with_next = True
        if spec.get("source"):
            sp = d.add_paragraph(); r = sp.add_run("%s: %s" % (self.L("source"), spec["source"])); r.font.size = Pt(8); r.font.color.rgb = rgbc(C.GREY); r.italic = True
        else:
            sp = d.add_paragraph(); sp.paragraph_format.space_after = Pt(4)

    def keyvalue(self, spec):
        d = self.d
        rows = spec["rows"]
        if spec.get("caption"):
            self.tab_n += 1
            cp = d.add_paragraph(style="Caption"); cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_after = Pt(3)
            cp.add_run("%s %d. %s" % (self.L("table"), self.tab_n, spec["caption"]))
        t = d.add_table(rows=len(rows), cols=2); table_no_borders(t); set_col_widths(t, [5.2, TEXT_W - 5.2])
        for i, (k, v) in enumerate(rows):
            a, b = t.cell(i, 0), t.cell(i, 1); shade(a, C.PALEBLUE)
            no_split(t.rows[i])
            r = a.paragraphs[0].add_run(str(k)); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = rgbc(C.NAVY)
            rich(b.paragraphs[0], v, size=9.5)
            for c_ in (a, b):
                cell_margins(c_, 55, 55, 100, 100); cell_borders(c_, bottom=(4, "D9DEE5"))
                c_.paragraphs[0].paragraph_format.space_after = Pt(0)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def figure(self, spec):
        d = self.d
        self.fig_n += 1
        w = min(spec.get("width_cm", 15), TEXT_W)
        with Image.open(spec["path"]) as im:
            ar = im.height / im.width
        if w * ar > 21:
            w = 21 / ar
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(spec["path"], width=Cm(w))
        cap = spec.get("caption")
        if cap:
            cp = d.add_paragraph(style="Caption"); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp.add_run("%s %d. %s" % (self.L("figure"), self.fig_n, cap))
            if spec.get("source"):
                cp.add_run(" (%s: %s)" % (self.L("source"), spec["source"]))

    def kpis(self, items):
        d = self.d
        n = len(items)
        t = d.add_table(rows=1, cols=n); table_no_borders(t); set_col_widths(t, [TEXT_W / n] * n)
        for j, it in enumerate(items):
            c = t.cell(0, j); shade(c, C.LIGHT); cell_margins(c, 120, 120, 120, 120)
            cell_borders(c, left=(18, "FFFFFF"), right=(18, "FFFFFF"), top=(18, C.BLUE))
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(it["value"])); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = rgbc(C.BLUE)
            if it.get("unit"):
                r = p.add_run(" " + it["unit"]); r.font.size = Pt(9); r.font.color.rgb = rgbc(C.GREY)
            p2 = c.add_paragraph(); p2.paragraph_format.space_after = Pt(0)
            r = p2.add_run(it["label"]); r.font.size = Pt(8.5); r.font.color.rgb = rgbc(C.GREY)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)

    def callout(self, spec):
        d = self.d
        kind = spec.get("kind", "info")
        col, bg = {"info": (C.BLUE, "EAF2FC"), "warn": (C.WARN, "FFF4DD"), "ok": (C.OK, "E3F4EA"), "bad": (C.BAD, "FBE6E3")}[kind]
        t = d.add_table(rows=1, cols=1); table_no_borders(t); set_col_widths(t, [TEXT_W])
        c = t.cell(0, 0); shade(c, bg); cell_margins(c, 110, 110, 160, 140); cell_borders(c, left=(36, col))
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        if spec.get("title"):
            r = p.add_run(spec["title"] + "  "); r.bold = True; r.font.color.rgb = rgbc(col)
        rich(p, spec["text"], size=10)
        d.add_paragraph().paragraph_format.space_after = Pt(2)

    # ---- carta (respuesta a requerimiento) ----
    def letterhead(self):
        d, m = self.d, self.meta
        p = d.add_paragraph()
        p.add_run().add_picture(C.logo("normawind-horizontal-positive.png"), width=Cm(5.2))
        d.add_paragraph()
        for ln in m.get("recipient", []):
            p = d.add_paragraph(); p.paragraph_format.space_after = Pt(0); r = p.add_run(ln); r.bold = (ln == m["recipient"][0])
        d.add_paragraph()
        place = m.get("place_date") or fmt_date(m.get("date"), self.lang)
        p = d.add_paragraph(place); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rows = m.get("reference_lines", [])
        if m.get("subject"):
            rows = [[{"en": "SUBJECT", "es": "ASUNTO", "ca": "ASSUMPTE"}[self.lang], m["subject"]]] + rows
        t = d.add_table(rows=len(rows), cols=2); table_no_borders(t); set_col_widths(t, [3.6, TEXT_W - 3.6])
        for i, (k, v) in enumerate(rows):
            a, b = t.cell(i, 0), t.cell(i, 1); shade(a, C.PALEBLUE)
            ra = a.paragraphs[0].add_run(k); ra.bold = True; ra.font.size = Pt(9.5); ra.font.color.rgb = rgbc(C.NAVY)
            rich(b.paragraphs[0], v, size=10)
            for c_ in (a, b):
                cell_margins(c_, 60, 60, 100, 100); cell_borders(c_, bottom=(4, "D9DEE5")); c_.paragraphs[0].paragraph_format.space_after = Pt(0)
        d.add_paragraph()

    # ---- montaje ----
    def build(self, update_toc_on_open=False):
        style = self.c.get("style") or C.report_types().get(self.meta.get("doc_type"), {}).get("style", "document")
        secs = self.c["sections"]
        num = style != "letter"
        if style == "letter":
            self.letterhead()  # la primera página no lleva cabecera: el membrete la sustituye
        else:
            self.cover(); self.control_page(); self.toc_page()
        nm = {id(n): k for k, lv, n in numbered(secs)}
        for s in secs:
            title = section_title(self.lang, s)
            self.heading(("%s.  %s" % (nm[id(s)], title)) if num else title, 1)
            for b in s["blocks"]:
                self.block(b)
            for sub in s.get("subsections", []) or []:
                st = sub.get("title") or section_title(self.lang, sub)
                self.heading(("%s.  %s" % (nm[id(sub)], st)) if num else st, 2)
                for b in sub["blocks"]:
                    self.block(b)
        if update_toc_on_open and style != "letter":
            st = self.d.settings.element
            e = OxmlElement("w:updateFields"); e.set(qn("w:val"), "true"); st.append(e)
        return self.d


def render(content, out_path, update_toc_on_open=False):
    doc = Doc(content).build(update_toc_on_open)
    doc.save(out_path)
    return out_path
